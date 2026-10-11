"""
engine/scores_xlsx.py — read SBI's Scores.xlsx directly (no copying into
Input_Docs). One sheet per case_id, three rows: DRC Category, CIBIL SCORE,
Profile Score. Already-structured tabular data — read with openpyxl, same
deterministic-join pattern as engine/policy_data.py. No LLM call needed
for this part; only the IIB screenshot images (agents/bureau.py) need
vision extraction.

Path resolution order:
  1. Explicit `path=` argument to `load_scores`.
  2. `SBI_SCORES_XLSX` environment variable.
  3. None — `lookup()` returns None for every case_id; SBI gating
     no-ops cleanly (same fallback behaviour as policy_data.py).
"""

from __future__ import annotations
import logging
import os
from pathlib import Path
from typing import Optional

from .reference_cache import source_revision

_loaded_revisions = {}
logger = logging.getLogger(__name__)

_INDEX: Optional[dict[str, dict]] = None
_LOADED_PATH: Optional[Path] = None


def _parse_cibil_score(val) -> tuple[Optional[int], str]:
    """A bureau CIBIL score is 300-900; some exports use a negative
    sentinel (observed: -1) for "no hit"/insufficient history rather
    than a literal 'no record found' string. Found via case SYNTHETIC_CASE
    (2026-10-01): an unflagged -1 was being compared as a real score,
    producing a misleading "CIBIL score -1 < 650" FAIL instead of the
    correct "no CIBIL record" WARN."""
    if val is None or str(val).strip().lower() in ("no record found", "no record", "", "-"):
        return None, "no_record_found"
    try:
        score = int(float(val))
    except (TypeError, ValueError):
        return None, "no_record_found"
    if score < 0:
        return None, "no_record_found"
    return score, "available"


def _resolve_path(path: Optional[str | Path] = None) -> Optional[Path]:
    if path:
        p = Path(path)
        return p if p.exists() else None
    env_p = os.environ.get("SBI_SCORES_XLSX")
    if env_p:
        p = Path(env_p)
        return p if p.exists() else None
    return None


def load_scores(path: Optional[str | Path] = None) -> dict[str, dict]:
    global _INDEX, _LOADED_PATH
    resolved = _resolve_path(path)
    if resolved is None:
        if _INDEX is None:
            logger.info("[scores_xlsx] no path resolved — SBI Scores.xlsx join disabled")
        _INDEX = {}
        _LOADED_PATH = None
        return _INDEX
    revision = source_revision(resolved)
    if (
        _INDEX is not None
        and _LOADED_PATH == resolved
        and _loaded_revisions.get("_LOADED_PATH") == revision
    ):
        return _INDEX

    try:
        import openpyxl

        wb = openpyxl.load_workbook(resolved, data_only=True)
    except Exception as e:
        logger.warning("[scores_xlsx] failed to load %s: %s", resolved, e)
        _INDEX, _LOADED_PATH = {}, None
        return _INDEX

    index: dict[str, dict] = {}
    for ws in wb.worksheets:
        case_id = ws.title.strip()
        if not case_id:
            continue
        rec: dict = {}
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            key = str(row[-2]).strip() if row[-2] is not None else ""
            val = row[-1]
            key_l = key.lower()
            if "drc" in key_l:
                rec["drc_category"] = str(val).strip() if val is not None else None
            elif "cibil" in key_l:
                rec["cibil_score"], rec["cibil_score_status"] = _parse_cibil_score(val)
            elif "profile" in key_l:
                rec["profile_score"] = str(val).strip() if val is not None else None
        if rec:
            index[case_id] = rec

    _INDEX = index
    _LOADED_PATH = resolved
    _loaded_revisions["_LOADED_PATH"] = revision
    logger.info("[scores_xlsx] indexed %d case(s) from %s", len(_INDEX), resolved.name)
    return _INDEX


def lookup_scores(case_id: str, path: Optional[str | Path] = None) -> Optional[dict]:
    """Return {drc_category, cibil_score, cibil_score_status, profile_score}
    for `case_id`, or None if not on file."""
    idx = load_scores(path)
    return idx.get(str(case_id).strip())


# ---------------------------------------------------------------------------
# Tabular variant — a later batch's Scores workbook arrived as ONE sheet
# with one ROW per case (columns PROPOSAL / WFL.CIBIL_SCORE / WFL.DRC_SCORE
# / PROFILE SCORE) instead of one sheet per case_id. Kept as a separate
# loader/env var (SBI_SCORES_XLSX_TABULAR) rather than reworking
# load_scores() above, so the original per-sheet batch keeps working
# unchanged — orchestrator/master.py tries both and merges.
# ---------------------------------------------------------------------------

_TAB_INDEX: Optional[dict[str, dict]] = None
_TAB_LOADED_PATH: Optional[Path] = None


def _resolve_tabular_path(path: Optional[str | Path] = None) -> Optional[Path]:
    if path:
        p = Path(path)
        return p if p.exists() else None
    env_p = os.environ.get("SBI_SCORES_XLSX_TABULAR")
    if env_p:
        p = Path(env_p)
        return p if p.exists() else None
    return None


def load_scores_tabular(path: Optional[str | Path] = None) -> dict[str, dict]:
    global _TAB_INDEX, _TAB_LOADED_PATH
    resolved = _resolve_tabular_path(path)
    if resolved is None:
        if _TAB_INDEX is None:
            logger.info(
                "[scores_xlsx] no tabular path resolved — SBI_SCORES_XLSX_TABULAR join disabled"
            )
        _TAB_INDEX = {}
        _TAB_LOADED_PATH = None
        return _TAB_INDEX
    revision = source_revision(resolved)
    if (
        _TAB_INDEX is not None
        and _TAB_LOADED_PATH == resolved
        and _loaded_revisions.get("_TAB_LOADED_PATH") == revision
        and _loaded_revisions.get("_LOADED_PATH") == revision
    ):
        return _TAB_INDEX

    try:
        import openpyxl

        wb = openpyxl.load_workbook(resolved, data_only=True)
    except Exception as e:
        logger.warning("[scores_xlsx] failed to load tabular %s: %s", resolved, e)
        _TAB_INDEX, _TAB_LOADED_PATH = {}, None
        return _TAB_INDEX

    index: dict[str, dict] = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
        except StopIteration:
            continue
        col = {name: i for i, name in enumerate(header)}
        prop_i = col.get("proposal")
        if prop_i is None:
            continue  # not this sheet's format
        cibil_i = next((i for n, i in col.items() if "cibil" in n), None)
        drc_i = next((i for n, i in col.items() if "drc" in n), None)
        profile_i = next((i for n, i in col.items() if "profile" in n), None)
        for row in rows:
            if not row or len(row) <= prop_i or not row[prop_i]:
                continue
            case_id = str(row[prop_i]).strip()
            if not case_id:
                continue
            rec: dict = {}
            if drc_i is not None and row[drc_i] not in (None, "", "-"):
                rec["drc_category"] = str(row[drc_i]).strip()
            if profile_i is not None and row[profile_i] not in (None, "", "-"):
                rec["profile_score"] = str(row[profile_i]).strip()
            if cibil_i is not None:
                rec["cibil_score"], rec["cibil_score_status"] = _parse_cibil_score(row[cibil_i])
            if rec:
                index[case_id] = rec

    _TAB_INDEX = index
    _TAB_LOADED_PATH = resolved
    _loaded_revisions["_TAB_LOADED_PATH"] = revision
    logger.info("[scores_xlsx] indexed %d case(s) from tabular %s", len(_TAB_INDEX), resolved.name)
    return _TAB_INDEX


def lookup_scores_tabular(case_id: str, path: Optional[str | Path] = None) -> Optional[dict]:
    idx = load_scores_tabular(path)
    return idx.get(str(case_id).strip())
